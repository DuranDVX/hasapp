"""Shared Claude client and one call helper."""
import base64
import io
import json
import os
from pathlib import Path

import anthropic

from . import config

_client = None


def client() -> anthropic.Anthropic:
    global _client
    if _client is None:
        if not os.getenv("ANTHROPIC_API_KEY"):
            raise RuntimeError("No ANTHROPIC_API_KEY. Add it to the .env file "
                               "and restart the server.")
        _client = anthropic.Anthropic()
    return _client


def call(system: str | list, content, *, schema: dict | None = None,
         model: str | None = None, effort: str | None = None,
         max_tokens: int = 16000):
    """Return parsed JSON when `schema` is given, else text. Streams, so long
    extractions do not hit HTTP timeouts."""
    output_config = {"effort": effort or config.EFFORT}
    if schema:
        output_config["format"] = {"type": "json_schema", "schema": schema}
    if isinstance(system, str):
        system = [{"type": "text", "text": system,
                   "cache_control": {"type": "ephemeral"}}]
    with client().beta.messages.stream(
        model=model or config.MODEL,
        max_tokens=max_tokens,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        output_config=output_config,
        system=system,
        messages=[{"role": "user", "content": content}],
    ) as stream:
        resp = stream.get_final_message()
    if resp.stop_reason == "refusal":
        raise RuntimeError("The AI declined this request.")
    if resp.stop_reason == "max_tokens":
        raise RuntimeError("The output was too long. Split the file and try again.")
    text = "".join(b.text for b in resp.content if b.type == "text")
    return json.loads(text) if schema else text


# ---- files -> content blocks ----

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".webp": "image/webp", ".gif": "image/gif"}


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode()


def _sheet_text(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"### Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                out.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(out)


def _docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(parts)


def file_blocks(name: str, data: bytes) -> list[dict]:
    """Turn an upload into content blocks Claude can read."""
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        return [{"type": "document", "source": {"type": "base64",
                 "media_type": "application/pdf", "data": _b64(data)}}]
    if ext in IMAGE_TYPES:
        return [{"type": "image", "source": {"type": "base64",
                 "media_type": IMAGE_TYPES[ext], "data": _b64(data)}}]
    if ext in (".xlsx", ".xlsm"):
        text = _sheet_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext in (".csv", ".txt", ".md", ".tsv"):
        text = data.decode("utf-8-sig", errors="replace")
    elif ext in (".heic", ".heif"):
        raise ValueError("HEIC photos are not supported. Set the camera to "
                         "'Most Compatible' or send a screenshot.")
    elif ext == ".xls":
        raise ValueError("Old .xls files are not supported. Save as .xlsx and upload again.")
    else:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    return [{"type": "text", "text": f"File: {name}\n\n{text}"}]
